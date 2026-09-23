"""Regression tests for the asset spreadsheet import.

The import writes many rows at once on behalf of a single click, so these pin
the two properties that matter most: a file is never partially applied, and a
row can only reference sites inside the uploader's own organization.
"""

import io

import pytest
from fastapi import status
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.auth import hash_password
from app.bulk_import import MAX_IMPORT_ROWS
from app.database import get_db
from app.main import app
from app.models import Agency, Asset, AssetType, CustodyType, Organization, User, UserRole, Warehouse

HEADERS = ["Asset ID", "Make / Model", "Asset Type", "Purchase Cost", "Location", "Custody Type", "Warehouse", "Agency", "Notes"]


@pytest.fixture
def client(db_session):
    def _override_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def org(db_session):
    """One importing org with a warehouse, plus a second org whose sites must stay invisible."""
    home = Organization(name="Import Org", slug="import-org", org_type="internal")
    other = Organization(name="Other Org", slug="other-import-org", org_type="internal")
    db_session.add_all([home, other])
    db_session.flush()

    manager = User(
        organization_id=home.id,
        email="manager@import.com",
        hashed_password=hash_password("password123"),
        full_name="Fleet Manager",
        role=UserRole.FLEET_MANAGER,
        is_active=True,
    )
    viewer = User(
        organization_id=home.id,
        email="viewer@import.com",
        hashed_password=hash_password("password123"),
        full_name="Read Only",
        role=UserRole.READ_ONLY,
        is_active=True,
    )
    home_warehouse = Warehouse(organization_id=home.id, name="Phoenix Depot")
    home_agency = Agency(organization_id=home.id, name="Tempe PD")
    other_warehouse = Warehouse(organization_id=other.id, name="Foreign Depot")
    db_session.add_all([manager, viewer, home_warehouse, home_agency, other_warehouse])
    db_session.commit()

    return {
        "home": home,
        "other": other,
        "warehouse": home_warehouse,
        "agency": home_agency,
        "other_warehouse": other_warehouse,
    }


def token_for(client, email="manager@import.com", password="password123"):
    response = client.post("/auth/login", json={"email": email, "password": password})
    assert response.status_code == status.HTTP_200_OK, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def make_xlsx(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers if headers is not None else HEADERS)
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_csv(rows, headers=None):
    lines = [",".join(headers if headers is not None else HEADERS)]
    lines.extend(",".join(str(cell) for cell in row) for row in rows)
    return "\n".join(lines).encode("utf-8")


def upload(client, headers, content, filename="assets.xlsx", commit=False):
    return client.post(
        f"/assets/import?commit={'true' if commit else 'false'}",
        files={"file": (filename, content, "application/octet-stream")},
        headers=headers,
    )


def row(vin="IMP-001", make_model="Flock Falcon", asset_type="ALPR Trailer", cost="1500", location="Yard", custody="", warehouse="", agency="", notes=""):
    return [vin, make_model, asset_type, cost, location, custody, warehouse, agency, notes]


class TestParsing:
    def test_xlsx_and_csv_produce_the_same_result(self, client, org):
        headers = token_for(client)
        rows = [row(vin="IMP-001"), row(vin="IMP-002", make_model="Ford F-150", asset_type="Fleet Vehicle")]

        from_xlsx = upload(client, headers, make_xlsx(rows)).json()
        from_csv = upload(client, headers, make_csv(rows), filename="assets.csv").json()

        assert from_xlsx["total_rows"] == from_csv["total_rows"] == 2
        assert from_xlsx["error_rows"] == from_csv["error_rows"] == 0

    def test_header_spelling_and_case_are_tolerated(self, client, org):
        headers = token_for(client)
        alt = ["vin", "MAKE_MODEL", "type", "Cost", "site", "custody", "warehouse", "agency", "notes"]

        response = upload(client, headers, make_xlsx([row()], headers=alt))

        assert response.json()["error_rows"] == 0

    def test_enum_values_accept_common_variants(self, client, org):
        headers = token_for(client)
        rows = [
            row(vin="IMP-001", asset_type="alpr_trailer", custody="warehouse"),
            row(vin="IMP-002", asset_type="Semi Truck", custody="in transit"),
        ]

        response = upload(client, headers, make_xlsx(rows))

        assert response.json()["error_rows"] == 0, response.json()["rows"]

    def test_currency_formatting_is_stripped(self, client, org):
        headers = token_for(client)

        response = upload(client, headers, make_xlsx([row(cost="$18,500.00")]), commit=True)

        assert response.json()["committed"] is True
        asset = client.get("/assets", headers=headers).json()[0]
        assert float(asset["initial_purchase_cost"]) == 18500.00

    def test_blank_rows_are_skipped(self, client, org):
        headers = token_for(client)
        rows = [row(vin="IMP-001"), ["", "", "", "", "", "", "", "", ""], row(vin="IMP-002")]

        assert upload(client, headers, make_xlsx(rows)).json()["total_rows"] == 2

    def test_reported_row_numbers_match_the_spreadsheet(self, client, org):
        headers = token_for(client)
        rows = [row(vin="IMP-001"), row(vin="IMP-002", cost="not-a-number")]

        results = upload(client, headers, make_xlsx(rows)).json()["rows"]

        # Header occupies row 1, so the bad row is row 3 in Excel's own numbering.
        assert [r["row_number"] for r in results] == [2, 3]
        assert results[1]["errors"]

    def test_missing_required_column_is_rejected(self, client, org):
        headers = token_for(client)
        short = ["Asset ID", "Make / Model", "Purchase Cost"]

        response = upload(client, headers, make_xlsx([["IMP-001", "Falcon", "100"]], headers=short))

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert "Asset Type" in response.json()["detail"]

    def test_two_columns_mapping_to_one_field_is_rejected(self, client, org):
        headers = token_for(client)
        dupe = ["Asset ID", "VIN", "Make / Model", "Asset Type", "Purchase Cost"]

        response = upload(client, headers, make_xlsx([["A", "B", "Falcon", "ALPR Trailer", "100"]], headers=dupe))

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert "both map to" in response.json()["detail"]

    def test_legacy_xls_gets_a_actionable_message(self, client, org):
        headers = token_for(client)

        response = upload(client, headers, b"\xd0\xcf\x11\xe0", filename="old.xls")

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert ".xlsx" in response.json()["detail"]

    def test_unsupported_extension_is_rejected(self, client, org):
        headers = token_for(client)

        response = upload(client, headers, b"whatever", filename="assets.pdf")

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert "Unsupported file type" in response.json()["detail"]

    def test_header_only_file_is_rejected(self, client, org):
        headers = token_for(client)

        response = upload(client, headers, make_xlsx([]))

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert "no data rows" in response.json()["detail"]

    def test_row_limit_is_enforced(self, client, org):
        headers = token_for(client)
        rows = [row(vin=f"IMP-{i:05d}") for i in range(MAX_IMPORT_ROWS + 1)]

        response = upload(client, headers, make_csv(rows), filename="big.csv")

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
        assert str(MAX_IMPORT_ROWS) in response.json()["detail"]


class TestValidation:
    def test_unknown_asset_type_reports_the_valid_options(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(asset_type="Spaceship")])).json()

        assert result["error_rows"] == 1
        assert "ALPR Trailer" in result["rows"][0]["errors"][0]

    def test_duplicate_vin_inside_the_file_points_at_the_other_row(self, client, org):
        headers = token_for(client)
        rows = [row(vin="SAME-1"), row(vin="same-1")]

        result = upload(client, headers, make_xlsx(rows)).json()

        assert result["error_rows"] == 1
        assert "row 2" in result["rows"][1]["errors"][0]

    def test_vin_already_in_the_fleet_is_rejected(self, client, org, db_session):
        db_session.add(
            Asset(
                organization_id=org["home"].id,
                vin="EXISTING-1",
                make_model="Already Here",
                initial_purchase_cost=100,
                current_location="Yard",
                asset_type=AssetType.ALPR_TRAILER,
            )
        )
        db_session.commit()
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(vin="EXISTING-1")])).json()

        assert "already exists" in result["rows"][0]["errors"][0]

    def test_missing_required_cells_are_listed_per_row(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(vin="", cost="")])).json()

        errors = " ".join(result["rows"][0]["errors"])
        assert "Asset ID is required" in errors
        assert "Purchase Cost is required" in errors

    def test_zero_and_negative_costs_are_rejected(self, client, org):
        headers = token_for(client)
        rows = [row(vin="IMP-001", cost="0"), row(vin="IMP-002", cost="-5")]

        result = upload(client, headers, make_xlsx(rows)).json()

        assert result["error_rows"] == 2
        assert all("greater than 0" in " ".join(r["errors"]) for r in result["rows"])

    def test_short_vin_is_rejected_by_the_same_rule_as_the_form(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(vin="AB")])).json()

        assert result["error_rows"] == 1

    def test_location_may_be_omitted_when_a_warehouse_is_named(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(location="", warehouse="Phoenix Depot")])).json()

        assert result["error_rows"] == 0, result["rows"]

    def test_location_required_without_a_site(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(location="")])).json()

        assert "Location is required" in " ".join(result["rows"][0]["errors"])

    def test_warehouse_and_agency_together_is_rejected(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(warehouse="Phoenix Depot", agency="Tempe PD")])).json()

        assert "not both" in " ".join(result["rows"][0]["errors"])

    def test_unknown_warehouse_name_is_rejected(self, client, org):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(warehouse="Nowhere Depot")])).json()

        assert "was not found" in " ".join(result["rows"][0]["errors"])


class TestTenantScoping:
    def test_another_orgs_warehouse_is_not_resolvable(self, client, org):
        """A name that exists only in a foreign org must read as 'not found'."""
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(warehouse="Foreign Depot")])).json()

        assert result["error_rows"] == 1
        assert "was not found" in " ".join(result["rows"][0]["errors"])

    def test_imported_assets_land_in_the_uploaders_organization(self, client, org, db_session):
        headers = token_for(client)

        upload(client, headers, make_xlsx([row(vin="SCOPE-1")]), commit=True)

        asset = db_session.query(Asset).filter(Asset.vin == "SCOPE-1").one()
        assert asset.organization_id == org["home"].id

    def test_template_only_lists_the_callers_own_sites(self, client, org):
        headers = token_for(client)

        response = client.get("/assets/import/template?format=csv", headers=headers)

        body = response.content.decode("utf-8-sig")
        assert "Phoenix Depot" in body
        assert "Foreign Depot" not in body


class TestCommit:
    def test_dry_run_creates_nothing(self, client, org, db_session):
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(vin="DRY-1")])).json()

        assert result["committed"] is False
        assert result["valid_rows"] == 1
        assert db_session.query(Asset).count() == 0

    def test_commit_creates_every_row(self, client, org, db_session):
        headers = token_for(client)
        rows = [row(vin=f"BULK-{i}") for i in range(5)]

        result = upload(client, headers, make_xlsx(rows), commit=True).json()

        assert result["committed"] is True
        assert result["created_count"] == 5
        assert db_session.query(Asset).count() == 5
        assert all(r["created_asset_id"] for r in result["rows"])

    def test_plate_columns_are_imported_and_normalized(self, client, org, db_session):
        headers = token_for(client)
        plate_headers = ["Asset ID", "Plate", "Plate State", "Make / Model", "Asset Type", "Purchase Cost", "Location"]
        rows = [
            ["PLT-1", "abc1234", "az", "Ford F-150", "Fleet Vehicle", "42750", "Yard"],
            ["PLT-2", "", "", "Flock Falcon", "ALPR Trailer", "18500", "Yard"],
        ]

        result = upload(client, headers, make_xlsx(rows, headers=plate_headers), commit=True).json()

        assert result["created_count"] == 2, result
        plated = db_session.query(Asset).filter_by(vin="PLT-1").one()
        assert plated.license_plate == "ABC1234"
        assert plated.license_plate_state == "AZ"
        assert db_session.query(Asset).filter_by(vin="PLT-2").one().license_plate is None

    def test_a_sheet_without_plate_columns_still_imports(self, client, org, db_session):
        """Plates are optional, so spreadsheets built before the column work as-is."""
        headers = token_for(client)

        result = upload(client, headers, make_xlsx([row(vin="NOPLATE-1")]), commit=True).json()

        assert result["created_count"] == 1, result
        assert db_session.query(Asset).one().license_plate is None

    def test_a_single_bad_row_blocks_the_whole_file(self, client, org, db_session):
        """All-or-nothing: 9 good rows must not land when row 10 is broken."""
        headers = token_for(client)
        rows = [row(vin=f"PART-{i}") for i in range(9)]
        rows.append(row(vin="PART-BAD", asset_type="Spaceship"))

        result = upload(client, headers, make_xlsx(rows), commit=True).json()

        assert result["committed"] is False
        assert result["error_rows"] == 1
        assert db_session.query(Asset).count() == 0

    def test_warehouse_row_links_the_asset_and_adopts_its_name(self, client, org, db_session):
        headers = token_for(client)

        upload(client, headers, make_xlsx([row(vin="WH-1", location="", warehouse="Phoenix Depot")]), commit=True)

        asset = db_session.query(Asset).filter(Asset.vin == "WH-1").one()
        assert asset.warehouse_id == org["warehouse"].id
        assert asset.current_location == "Phoenix Depot"

    def test_custody_drives_operational_status_like_the_single_asset_form(self, client, org, db_session):
        headers = token_for(client)
        rows = [
            row(vin="CUST-1", custody="Warehouse Depot", warehouse="Phoenix Depot", location=""),
            row(vin="CUST-2", custody="Customer / LE Agency", agency="Tempe PD", location=""),
        ]

        upload(client, headers, make_xlsx(rows), commit=True)

        by_vin = {a.vin: a for a in db_session.query(Asset).all()}
        assert by_vin["CUST-1"].operational_status.value == "available"
        assert by_vin["CUST-2"].operational_status.value == "deployed"
        assert by_vin["CUST-2"].current_custody_type == CustodyType.CUSTOMER_AGENCY

    def test_vin_is_normalized_to_uppercase(self, client, org, db_session):
        headers = token_for(client)

        upload(client, headers, make_xlsx([row(vin="lower-case-vin")]), commit=True)

        assert db_session.query(Asset).one().vin == "LOWER-CASE-VIN"


class TestTemplate:
    def test_template_round_trips_through_the_importer(self, client, org):
        """The sheet we hand out must validate cleanly once the example rows are filled in."""
        headers = token_for(client)

        template = client.get("/assets/import/template", headers=headers).content
        workbook = load_workbook(io.BytesIO(template))
        sheet = workbook.active
        # Row 2 is a worked example; uniquify the ID and upload it back.
        sheet.cell(row=2, column=1).value = "ROUNDTRIP-1"
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = upload(client, headers, buffer.getvalue()).json()

        assert result["error_rows"] == 0, result["rows"]

    def test_reference_comment_lines_do_not_become_assets(self, client, org):
        headers = token_for(client)

        template = client.get("/assets/import/template", headers=headers).content
        result = upload(client, headers, template).json()

        # Only the two example rows count; the '# Reference' block is not data.
        assert result["total_rows"] == 2

    def test_csv_template_is_excel_readable(self, client, org):
        headers = token_for(client)

        response = client.get("/assets/import/template?format=csv", headers=headers)

        assert response.content.startswith(b"\xef\xbb\xbf")  # BOM so Excel detects UTF-8
        assert "attachment" in response.headers["content-disposition"]

    def test_unknown_format_is_rejected(self, client, org):
        headers = token_for(client)

        response = client.get("/assets/import/template?format=pdf", headers=headers)

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY


class TestPermissions:
    def test_read_only_user_cannot_import(self, client, org):
        headers = token_for(client, email="viewer@import.com")

        response = upload(client, headers, make_xlsx([row()]), commit=True)

        assert response.status_code == status.HTTP_403_FORBIDDEN

    def test_read_only_user_cannot_download_the_template(self, client, org):
        headers = token_for(client, email="viewer@import.com")

        assert client.get("/assets/import/template", headers=headers).status_code == status.HTTP_403_FORBIDDEN

    def test_anonymous_upload_is_rejected(self, client, org):
        response = client.post(
            "/assets/import",
            files={"file": ("assets.xlsx", make_xlsx([row()]), "application/octet-stream")},
        )

        assert response.status_code == status.HTTP_401_UNAUTHORIZED
