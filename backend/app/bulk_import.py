"""Spreadsheet import for assets.

Accepts .xlsx or .csv and checks the whole file before writing anything, so a
bad cell in row 40 cannot leave rows 1-39 half-imported. Callers validate first
(dry run), show the user what is wrong, then commit once the file is clean.

Rows are turned into AssetCreate payloads and handed to ops.create_asset(), so
the import path cannot accept an asset the single-asset form would reject.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from uuid import UUID

from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Agency, Asset, AssetType, CustodyType, Warehouse
from app.schemas import AssetCreate

# A spreadsheet is a person-sized tool; these caps keep one upload from pinning
# the worker while still clearing any realistic fleet onboarding batch.
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

CSV_SUFFIXES = (".csv",)
EXCEL_SUFFIXES = (".xlsx", ".xlsm")


def _key(value: object) -> str:
    """Collapse a header or enum label to a comparison key ('Make / Model' -> 'makemodel')."""
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


@dataclass(frozen=True)
class Column:
    name: str
    label: str
    required: bool
    help: str
    aliases: tuple[str, ...] = ()

    def matches(self, header_key: str) -> bool:
        return header_key == _key(self.label) or header_key in {_key(a) for a in self.aliases}


COLUMNS: tuple[Column, ...] = (
    Column(
        "vin",
        "Asset ID",
        True,
        "Unique ID or VIN, 4-32 characters.",
        ("vin", "asset_id", "asset id / vin", "asset tag", "serial", "serial number"),
    ),
    Column(
        "license_plate",
        "Plate",
        False,
        "Registration plate, up to 16 characters. Leave blank for unplated units.",
        ("license_plate", "license plate", "plate number", "plate #", "tag", "tag number", "registration"),
    ),
    Column(
        "license_plate_state",
        "Plate State",
        False,
        "Two-letter state or province code for the plate, e.g. AZ.",
        ("license_plate_state", "plate state", "state", "issuing state", "reg state"),
    ),
    Column(
        "make_model",
        "Make / Model",
        True,
        "Manufacturer and model, e.g. 'Flock Safety Falcon'.",
        ("make_model", "make and model", "model", "description"),
    ),
    Column(
        "asset_type",
        "Asset Type",
        True,
        "One of: " + ", ".join(t.value for t in AssetType),
        ("asset_type", "type", "category"),
    ),
    Column(
        "initial_purchase_cost",
        "Purchase Cost",
        True,
        "Number greater than 0. '$' and ',' are ignored.",
        ("initial_purchase_cost", "cost", "purchase price", "price", "value"),
    ),
    Column(
        "current_location",
        "Location",
        False,
        "Free-text site. Optional when Warehouse or Agency is filled in.",
        ("current_location", "initial location", "site"),
    ),
    Column(
        "custody_type",
        "Custody Type",
        False,
        "One of: " + ", ".join(c.value for c in CustodyType) + ". Defaults to Warehouse Depot.",
        ("custody_type", "current_custody_type", "custody"),
    ),
    Column(
        "warehouse",
        "Warehouse",
        False,
        "Name of an existing warehouse. Sets the location automatically.",
        ("warehouse name", "depot"),
    ),
    Column(
        "agency",
        "Agency",
        False,
        "Name of an existing agency. Sets the location automatically.",
        ("agency name", "customer"),
    ),
    Column("notes", "Notes", False, "Optional free text.", ("note", "comment", "comments")),
)

COLUMNS_BY_NAME = {column.name: column for column in COLUMNS}

# Spelling variants we accept for enum cells, beyond the exact stored labels.
ASSET_TYPE_ALIASES = {
    "alpr": AssetType.ALPR_TRAILER,
    "trailer": AssetType.ALPR_TRAILER,
    "alprtrailer": AssetType.ALPR_TRAILER,
    "semi": AssetType.SEMI_TRUCK,
    "truck": AssetType.SEMI_TRUCK,
    "semitruck": AssetType.SEMI_TRUCK,
    "vehicle": AssetType.FLEET_VEHICLE,
    "fleetvehicle": AssetType.FLEET_VEHICLE,
    "car": AssetType.FLEET_VEHICLE,
}

CUSTODY_TYPE_ALIASES = {
    "warehouse": CustodyType.WAREHOUSE_DEPOT,
    "depot": CustodyType.WAREHOUSE_DEPOT,
    "warehousedepot": CustodyType.WAREHOUSE_DEPOT,
    "storage": CustodyType.WAREHOUSE_DEPOT,
    "transit": CustodyType.IN_TRANSIT,
    "intransit": CustodyType.IN_TRANSIT,
    "shipping": CustodyType.IN_TRANSIT,
    "customer": CustodyType.CUSTOMER_AGENCY,
    "agency": CustodyType.CUSTOMER_AGENCY,
    "customeragency": CustodyType.CUSTOMER_AGENCY,
    "customerleagency": CustodyType.CUSTOMER_AGENCY,
    "deployed": CustodyType.CUSTOMER_AGENCY,
}


class ImportFileError(Exception):
    """The file could not be read far enough to produce per-row feedback."""


@dataclass
class ParsedRow:
    row_number: int  # 1-based spreadsheet row, so users can find it in Excel
    values: dict[str, str]


@dataclass
class PreparedRow:
    row_number: int
    vin: str | None
    make_model: str | None
    errors: list[str] = field(default_factory=list)
    payload: AssetCreate | None = None


# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # openpyxl hands back 12345.0 for a plain integer cell; don't show ".0".
        return str(int(value))
    return str(value).strip()


def _rows_from_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("Could not read the file as text. Save it as UTF-8 CSV or .xlsx and try again.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[_cell_text(cell) for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _rows_from_excel(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - dependency is declared
        raise ImportFileError("Excel support is unavailable on the server. Upload a .csv instead.") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(
            "That file could not be opened as a workbook. If it is an older .xls file, re-save it as .xlsx."
        ) from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_rows(filename: str, data: bytes) -> list[list[str]]:
    """Return raw grid cells from a .csv or .xlsx upload."""
    if not data:
        raise ImportFileError("The uploaded file is empty.")
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFileError(f"File is larger than {MAX_IMPORT_BYTES // (1024 * 1024)} MB.")

    lowered = (filename or "").lower()
    if lowered.endswith(CSV_SUFFIXES):
        return _rows_from_csv(data)
    if lowered.endswith(EXCEL_SUFFIXES):
        return _rows_from_excel(data)
    if lowered.endswith(".xls"):
        raise ImportFileError("The old .xls format is not supported. Open it in Excel and save as .xlsx.")
    raise ImportFileError("Unsupported file type. Upload a .xlsx or .csv file.")


def parse_sheet(filename: str, data: bytes) -> list[ParsedRow]:
    """Map the grid onto our column names using the header row."""
    grid = read_rows(filename, data)

    header_index = next((i for i, row in enumerate(grid) if any(cell for cell in row)), None)
    if header_index is None:
        raise ImportFileError("The file has no rows.")

    header = grid[header_index]
    mapping: dict[int, str] = {}
    claimed: dict[str, str] = {}
    for position, raw_header in enumerate(header):
        header_key = _key(raw_header)
        if not header_key:
            continue
        column = next((c for c in COLUMNS if c.matches(header_key)), None)
        if column is None:
            continue  # extra columns are ignored, not an error
        if column.name in claimed:
            raise ImportFileError(
                f"Columns '{claimed[column.name]}' and '{raw_header}' both map to {column.label}. "
                "Remove or rename one of them."
            )
        claimed[column.name] = str(raw_header)
        mapping[position] = column.name

    missing = [c.label for c in COLUMNS if c.required and c.name not in claimed]
    if missing:
        raise ImportFileError(
            "The file is missing required column(s): "
            + ", ".join(missing)
            + ". Download the template to see the expected headers."
        )

    rows: list[ParsedRow] = []
    for offset, raw_row in enumerate(grid[header_index + 1 :], start=header_index + 2):
        values = {name: "" for name in COLUMNS_BY_NAME}
        for position, name in mapping.items():
            if position < len(raw_row):
                values[name] = raw_row[position]
        if not any(values.values()):
            continue  # blank spacer row
        first_cell = next((cell for cell in raw_row if cell), "")
        if first_cell.startswith("#"):
            continue  # reference/comment line, e.g. the block in our template
        rows.append(ParsedRow(row_number=offset, values=values))

    if not rows:
        raise ImportFileError("The file has headers but no data rows.")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportFileError(f"File has {len(rows)} rows; the limit is {MAX_IMPORT_ROWS} per upload.")
    return rows


# --------------------------------------------------------------------------
# Validating rows
# --------------------------------------------------------------------------


def _parse_enum(raw: str, enum_cls, aliases: dict) -> object | None:
    key = _key(raw)
    for member in enum_cls:
        if _key(member.value) == key or _key(member.name) == key:
            return member
    return aliases.get(key)


def _parse_money(raw: str) -> Decimal | None:
    cleaned = re.sub(r"[^0-9.\-]", "", raw)
    if not cleaned:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _directory_lookup(db: Session, model, organization_id: UUID) -> dict[str, list[UUID]]:
    """Name key -> ids, scoped to the caller's organization."""
    lookup: dict[str, list[UUID]] = {}
    rows = db.execute(
        select(model.id, model.name).where(
            model.organization_id == organization_id, model.is_archived.is_(False)
        )
    ).all()
    for row_id, name in rows:
        lookup.setdefault(_key(name), []).append(row_id)
    return lookup


def _resolve_directory(raw: str, lookup: dict[str, list[UUID]], label: str, errors: list[str]) -> UUID | None:
    matches = lookup.get(_key(raw))
    if not matches:
        errors.append(f"{label} '{raw}' was not found. Create it first or clear the cell.")
        return None
    if len(matches) > 1:
        errors.append(f"{label} '{raw}' matches more than one record. Rename them so the name is unique.")
        return None
    return matches[0]


def _friendly_validation_errors(exc: ValidationError) -> list[str]:
    messages = []
    for error in exc.errors():
        field_name = str(error["loc"][0]) if error.get("loc") else ""
        column = COLUMNS_BY_NAME.get(field_name)
        label = column.label if column else (field_name or "Row")
        messages.append(f"{label}: {error['msg']}")
    return messages


def prepare_rows(db: Session, rows: list[ParsedRow], organization_id: UUID) -> list[PreparedRow]:
    """Turn parsed rows into AssetCreate payloads, collecting per-row problems."""
    warehouses = _directory_lookup(db, Warehouse, organization_id)
    agencies = _directory_lookup(db, Agency, organization_id)

    wanted_vins = {
        row.values["vin"].strip().lower() for row in rows if row.values["vin"].strip()
    }
    taken_vins: set[str] = set()
    if wanted_vins:
        taken_vins = {
            vin.lower()
            for vin in db.scalars(
                select(Asset.vin).where(func.lower(Asset.vin).in_(wanted_vins))
            ).all()
        }

    seen_vins: dict[str, int] = {}
    prepared: list[PreparedRow] = []

    for row in rows:
        values = row.values
        errors: list[str] = []
        vin = values["vin"].strip().upper()
        make_model = values["make_model"].strip()
        result = PreparedRow(row_number=row.row_number, vin=vin or None, make_model=make_model or None)

        if not vin:
            errors.append("Asset ID is required.")
        elif vin.lower() in taken_vins:
            errors.append(f"Asset ID '{vin}' already exists in the fleet.")
        elif vin.lower() in seen_vins:
            errors.append(f"Asset ID '{vin}' is also used on row {seen_vins[vin.lower()]} of this file.")
        else:
            seen_vins[vin.lower()] = row.row_number

        asset_type = None
        if not values["asset_type"].strip():
            errors.append("Asset Type is required.")
        else:
            asset_type = _parse_enum(values["asset_type"], AssetType, ASSET_TYPE_ALIASES)
            if asset_type is None:
                errors.append(
                    f"Asset Type '{values['asset_type'].strip()}' is not recognized. "
                    "Use one of: " + ", ".join(t.value for t in AssetType) + "."
                )

        custody_type = CustodyType.WAREHOUSE_DEPOT
        if values["custody_type"].strip():
            parsed_custody = _parse_enum(values["custody_type"], CustodyType, CUSTODY_TYPE_ALIASES)
            if parsed_custody is None:
                errors.append(
                    f"Custody Type '{values['custody_type'].strip()}' is not recognized. "
                    "Use one of: " + ", ".join(c.value for c in CustodyType) + "."
                )
            else:
                custody_type = parsed_custody

        cost = None
        if not values["initial_purchase_cost"].strip():
            errors.append("Purchase Cost is required.")
        else:
            cost = _parse_money(values["initial_purchase_cost"])
            if cost is None:
                errors.append(f"Purchase Cost '{values['initial_purchase_cost'].strip()}' is not a number.")
            elif cost <= 0:
                errors.append("Purchase Cost must be greater than 0.")

        warehouse_id = None
        agency_id = None
        warehouse_name = values["warehouse"].strip()
        agency_name = values["agency"].strip()
        if warehouse_name and agency_name:
            errors.append("Fill in either Warehouse or Agency, not both.")
        elif warehouse_name:
            warehouse_id = _resolve_directory(warehouse_name, warehouses, "Warehouse", errors)
        elif agency_name:
            agency_id = _resolve_directory(agency_name, agencies, "Agency", errors)

        # create_asset() overwrites location with the linked site's name, so the
        # column is only required when the row names no warehouse or agency.
        location = values["current_location"].strip()
        if not location:
            if warehouse_name or agency_name:
                location = warehouse_name or agency_name
            else:
                errors.append("Location is required when no Warehouse or Agency is given.")

        if errors:
            result.errors = errors
            prepared.append(result)
            continue

        try:
            result.payload = AssetCreate(
                vin=vin,
                license_plate=values["license_plate"].strip() or None,
                license_plate_state=values["license_plate_state"].strip() or None,
                make_model=make_model,
                initial_purchase_cost=cost,
                current_location=location,
                asset_type=asset_type,
                current_custody_type=custody_type,
                warehouse_id=warehouse_id,
                agency_id=agency_id,
                notes=values["notes"].strip() or None,
            )
        except ValidationError as exc:
            result.errors = _friendly_validation_errors(exc)
        prepared.append(result)

    return prepared


# --------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------

EXAMPLE_ROWS = (
    ("ALPR-1001", "", "", "Flock Safety Falcon", "ALPR Trailer", "18500", "Main Yard", "Warehouse Depot", "", "", "Delete this example row"),
    ("VEH-2043", "ABC1234", "AZ", "Ford F-150", "Fleet Vehicle", "42750.00", "Phoenix Yard", "Warehouse Depot", "", "", ""),
)


def template_rows(warehouse_names: list[str], agency_names: list[str]) -> list[list[str]]:
    """Header + examples + an inline reference block, shared by the CSV and Excel templates."""
    rows: list[list[str]] = [[column.label for column in COLUMNS]]
    rows.extend(list(example) for example in EXAMPLE_ROWS)
    rows.append([])
    rows.append(["# Reference - delete these lines before uploading"])
    for column in COLUMNS:
        flag = "required" if column.required else "optional"
        rows.append([f"# {column.label} ({flag}): {column.help}"])
    if warehouse_names:
        rows.append(["# Your warehouses: " + ", ".join(warehouse_names)])
    if agency_names:
        rows.append(["# Your agencies: " + ", ".join(agency_names)])
    return rows


def template_csv(warehouse_names: list[str], agency_names: list[str]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(template_rows(warehouse_names, agency_names))
    # Excel needs the BOM to open a UTF-8 CSV without mangling accents.
    return buffer.getvalue().encode("utf-8-sig")


def template_xlsx(warehouse_names: list[str], agency_names: list[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assets"
    for row in template_rows(warehouse_names, agency_names):
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for position, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=position).column_letter].width = max(
            14, len(column.label) + 4
        )
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
